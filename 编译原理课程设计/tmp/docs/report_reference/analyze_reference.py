from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pdfplumber
from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def twips_to_cm(value):
    return None if value is None else round(int(value) / 1440 * 2.54, 3)


def emu_to_cm(value):
    return None if value is None else round(int(value) / 914400 * 2.54, 3)


def qn(name: str) -> str:
    prefix, local = name.split(":", 1)
    return f"{{{NS[prefix]}}}{local}"


def attr(el, name):
    return None if el is None else el.get(qn(name))


def bool_prop(parent, tag):
    el = parent.find(f"w:{tag}", NS) if parent is not None else None
    if el is None:
        return None
    val = attr(el, "w:val")
    return val not in {"0", "false", "off"}


def style_to_dict(style):
    el = style._element
    ppr = el.find("w:pPr", NS)
    rpr = el.find("w:rPr", NS)
    fonts = rpr.find("w:rFonts", NS) if rpr is not None else None
    sz = rpr.find("w:sz", NS) if rpr is not None else None
    color = rpr.find("w:color", NS) if rpr is not None else None
    jc = ppr.find("w:jc", NS) if ppr is not None else None
    spacing = ppr.find("w:spacing", NS) if ppr is not None else None
    ind = ppr.find("w:ind", NS) if ppr is not None else None
    tabs = ppr.findall("w:tabs/w:tab", NS) if ppr is not None else []
    based = el.find("w:basedOn", NS)
    return {
        "name": style.name,
        "type": str(style.type),
        "based_on": attr(based, "w:val"),
        "font_ascii": attr(fonts, "w:ascii"),
        "font_hansi": attr(fonts, "w:hAnsi"),
        "font_eastasia": attr(fonts, "w:eastAsia"),
        "font_cs": attr(fonts, "w:cs"),
        "size_pt": None if sz is None or attr(sz, "w:val") is None else int(attr(sz, "w:val")) / 2,
        "bold": bool_prop(rpr, "b"),
        "italic": bool_prop(rpr, "i"),
        "color": attr(color, "w:val"),
        "alignment": attr(jc, "w:val"),
        "before_pt": None if spacing is None or attr(spacing, "w:before") is None else int(attr(spacing, "w:before")) / 20,
        "after_pt": None if spacing is None or attr(spacing, "w:after") is None else int(attr(spacing, "w:after")) / 20,
        "line": attr(spacing, "w:line"),
        "line_rule": attr(spacing, "w:lineRule"),
        "left_cm": twips_to_cm(attr(ind, "w:left")),
        "right_cm": twips_to_cm(attr(ind, "w:right")),
        "first_line_cm": twips_to_cm(attr(ind, "w:firstLine")),
        "hanging_cm": twips_to_cm(attr(ind, "w:hanging")),
        "keep_next": bool_prop(ppr, "keepNext"),
        "keep_lines": bool_prop(ppr, "keepLines"),
        "page_break_before": bool_prop(ppr, "pageBreakBefore"),
        "tabs": [
            {"val": attr(t, "w:val"), "pos_cm": twips_to_cm(attr(t, "w:pos")), "leader": attr(t, "w:leader")}
            for t in tabs
        ],
    }


def run_signature(run):
    el = run._element
    rpr = el.find("w:rPr", NS)
    fonts = rpr.find("w:rFonts", NS) if rpr is not None else None
    sz = rpr.find("w:sz", NS) if rpr is not None else None
    color = rpr.find("w:color", NS) if rpr is not None else None
    highlight = rpr.find("w:highlight", NS) if rpr is not None else None
    shd = rpr.find("w:shd", NS) if rpr is not None else None
    return (
        attr(fonts, "w:ascii"),
        attr(fonts, "w:eastAsia"),
        None if sz is None or attr(sz, "w:val") is None else int(attr(sz, "w:val")) / 2,
        bool_prop(rpr, "b"),
        bool_prop(rpr, "i"),
        attr(color, "w:val"),
        attr(highlight, "w:val"),
        attr(shd, "w:fill"),
    )


def para_signature(p):
    ppr = p._p.find("w:pPr", NS)
    jc = ppr.find("w:jc", NS) if ppr is not None else None
    spacing = ppr.find("w:spacing", NS) if ppr is not None else None
    ind = ppr.find("w:ind", NS) if ppr is not None else None
    numpr = ppr.find("w:numPr", NS) if ppr is not None else None
    ilvl = numpr.find("w:ilvl", NS) if numpr is not None else None
    numid = numpr.find("w:numId", NS) if numpr is not None else None
    return {
        "alignment": attr(jc, "w:val"),
        "before_pt": None if spacing is None or attr(spacing, "w:before") is None else int(attr(spacing, "w:before")) / 20,
        "after_pt": None if spacing is None or attr(spacing, "w:after") is None else int(attr(spacing, "w:after")) / 20,
        "line": attr(spacing, "w:line"),
        "line_rule": attr(spacing, "w:lineRule"),
        "left_cm": twips_to_cm(attr(ind, "w:left")),
        "right_cm": twips_to_cm(attr(ind, "w:right")),
        "first_line_cm": twips_to_cm(attr(ind, "w:firstLine")),
        "hanging_cm": twips_to_cm(attr(ind, "w:hanging")),
        "num_id": attr(numid, "w:val"),
        "num_level": attr(ilvl, "w:val"),
    }


def table_info(table, index):
    tbl = table._tbl
    tblpr = tbl.find("w:tblPr", NS)
    tblw = tblpr.find("w:tblW", NS) if tblpr is not None else None
    tblind = tblpr.find("w:tblInd", NS) if tblpr is not None else None
    align = tblpr.find("w:jc", NS) if tblpr is not None else None
    borders = tblpr.find("w:tblBorders", NS) if tblpr is not None else None
    grid = tbl.find("w:tblGrid", NS)
    widths = [twips_to_cm(attr(c, "w:w")) for c in grid.findall("w:gridCol", NS)] if grid is not None else []
    fills = Counter()
    cell_margins = []
    verticals = Counter()
    for row in table.rows:
        for cell in row.cells:
            tcpr = cell._tc.find("w:tcPr", NS)
            shd = tcpr.find("w:shd", NS) if tcpr is not None else None
            fills[attr(shd, "w:fill") or "none"] += 1
            va = tcpr.find("w:vAlign", NS) if tcpr is not None else None
            verticals[attr(va, "w:val") or "default"] += 1
            tcm = tcpr.find("w:tcMar", NS) if tcpr is not None else None
            if tcm is not None:
                cell_margins.append({side: twips_to_cm(attr(tcm.find(f"w:{side}", NS), "w:w")) for side in ["top", "left", "bottom", "right"]})
    top_text = [[re.sub(r"\s+", " ", c.text).strip()[:120] for c in row.cells] for row in table.rows[:3]]
    border_desc = {}
    if borders is not None:
        for child in borders:
            border_desc[etree.QName(child).localname] = {
                "val": attr(child, "w:val"),
                "size_eighth_pt": attr(child, "w:sz"),
                "color": attr(child, "w:color"),
            }
    return {
        "index": index,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "width_cm": twips_to_cm(attr(tblw, "w:w")),
        "width_type": attr(tblw, "w:type"),
        "indent_cm": twips_to_cm(attr(tblind, "w:w")),
        "alignment": attr(align, "w:val"),
        "grid_cm": widths,
        "style": table.style.name if table.style else None,
        "fills": fills,
        "vertical_alignment": verticals,
        "cell_margins_examples": cell_margins[:5],
        "borders": border_desc,
        "sample": top_text,
    }


def extract_pdf(pdf_path: Path):
    pages = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            pages.append({
                "page": i,
                "width_pt": round(page.width, 2),
                "height_pt": round(page.height, 2),
                "text": text,
            })
    return pages


def main(docx_path: str, pdf_path: str, out_path: str):
    docx_path = Path(docx_path)
    pdf_path = Path(pdf_path)
    doc = Document(docx_path)
    style_counts = Counter(p.style.name if p.style else "" for p in doc.paragraphs)
    nonempty = []
    run_sigs = Counter()
    para_sigs = Counter()
    for i, p in enumerate(doc.paragraphs):
        text = re.sub(r"\s+", " ", p.text).strip()
        if text:
            nonempty.append({
                "index": i,
                "style": p.style.name if p.style else None,
                "text": text,
                "para_format": para_signature(p),
                "runs": [
                    {"text": re.sub(r"\s+", " ", r.text).strip()[:120], "format": run_signature(r)}
                    for r in p.runs if r.text.strip()
                ],
            })
        para_sigs[json.dumps(para_signature(p), ensure_ascii=False, sort_keys=True)] += 1
        for r in p.runs:
            if r.text.strip():
                run_sigs[str(run_signature(r))] += len(r.text)

    sects = []
    for i, s in enumerate(doc.sections, 1):
        sects.append({
            "index": i,
            "page_width_cm": emu_to_cm(s.page_width),
            "page_height_cm": emu_to_cm(s.page_height),
            "left_margin_cm": emu_to_cm(s.left_margin),
            "right_margin_cm": emu_to_cm(s.right_margin),
            "top_margin_cm": emu_to_cm(s.top_margin),
            "bottom_margin_cm": emu_to_cm(s.bottom_margin),
            "header_distance_cm": emu_to_cm(s.header_distance),
            "footer_distance_cm": emu_to_cm(s.footer_distance),
            "different_first_page_header_footer": s.different_first_page_header_footer,
            "header_text": [p.text for p in s.header.paragraphs],
            "footer_text": [p.text for p in s.footer.paragraphs],
        })

    styles = [style_to_dict(s) for s in doc.styles if s.type in (1, 2) and (style_counts.get(s.name, 0) or s.name in {"Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "TOC 1", "TOC 2", "TOC 3"})]
    data = {
        "docx": str(docx_path),
        "pdf": str(pdf_path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "style_counts": style_counts,
        "sections": sects,
        "styles": styles,
        "run_signature_weighted_chars": run_sigs.most_common(50),
        "paragraph_format_counts": para_sigs.most_common(30),
        "nonempty_paragraphs": nonempty,
        "tables": [table_info(t, i) for i, t in enumerate(doc.tables, 1)],
        "pdf_pages": extract_pdf(pdf_path),
    }
    Path(out_path).write_text(json.dumps(data, ensure_ascii=False, indent=2, default=lambda o: dict(o)), encoding="utf-8")


if __name__ == "__main__":
    main(*sys.argv[1:4])
